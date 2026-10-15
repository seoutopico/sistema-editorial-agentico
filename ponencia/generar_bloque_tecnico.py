# -*- coding: utf-8 -*-
"""Bloque técnico de «De chat a sistemas» (ConIA Madrid 2026), versión 2.

Patrón visual tomado de las presentaciones de Aina (ConIA y MujeresEnSEO):
título grande en Inter Bold, un subtítulo corto en coral, tres o cuatro líneas de
texto a la izquierda y, a la derecha, una ilustración técnica de las suyas o una
captura real de un archivo del repositorio con un recuadro rojo. Sin animaciones.

Requiere haber ejecutado antes `python ponencia/recursos.py` (capturas).
Salida: ponencia/ConIA_bloque_tecnico.pptx
"""
from pathlib import Path

from PIL import Image
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

AQUI = Path(__file__).parent
CAP = AQUI / "capturas"
ILU = AQUI / "ilustraciones"
LOGO = AQUI / "logo_thinkhack.png"
SALIDA = AQUI / "ConIA_bloque_tecnico.pptx"

CREMA = RGBColor(0xEA, 0xE7, 0xDC)
CORAL = RGBColor(0xFD, 0x65, 0x34)
TINTA = RGBColor(0x1F, 0x29, 0x37)
GRIS = RGBColor(0x55, 0x5B, 0x61)
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)
FUENTE = "Inter"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(10), Inches(5.625)
BLANK = prs.slide_layouts[6]


# ------------------------------------------------------------------ utilidades
def _run(p, tx, tam, negrita=False, color=TINTA, nombre=FUENTE):
    r = p.add_run()
    r.text = tx
    r.font.name = nombre
    r.font.size = Pt(tam)
    r.font.bold = negrita
    r.font.color.rgb = color
    return r


def nueva(notas):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = CREMA
    s.shapes.add_picture(str(LOGO), Inches(0.28), Inches(5.1), height=Inches(0.3))
    tb = s.shapes.add_textbox(Inches(0.62), Inches(5.1), Inches(3), Inches(0.3))
    tb.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    _run(tb.text_frame.paragraphs[0], "Think&Hack by ainalluna.com", 8)
    s.notes_slide.notes_text_frame.text = notas.strip()
    return s


def titulo(s, partes, y=Inches(0.45), tam=None, w=Inches(9.2)):
    """partes: lista de (texto, en_coral). Coral = fondo coral con texto blanco, como «sistemas».
    El tamaño baja con la longitud para que el título ocupe una sola línea."""
    n = sum(len(t) for t, _ in partes)
    if tam is None:
        tam = 40 if n <= 26 else 36 if n <= 30 else 30
    tb = s.shapes.add_textbox(Inches(0.4), y, w, Inches(1.2))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.line_spacing = 1.05
    for tx, coral in partes:
        r = _run(p, tx, tam, True, BLANCO if coral else TINTA)
        if coral:
            rPr = r._r.get_or_add_rPr()
            hl = etree.Element(qn("a:highlight"))
            etree.SubElement(hl, qn("a:srgbClr")).set("val", "FD6534")
            latin = rPr.find(qn("a:latin"))
            rPr.insert(list(rPr).index(latin), hl)
    return tb


def subtitulo(s, tx, y=Inches(1.75), x=Inches(0.4), w=Inches(3.7)):
    tb = s.shapes.add_textbox(x, y, w, Inches(0.4))
    tb.text_frame.word_wrap = True
    _run(tb.text_frame.paragraphs[0], tx, 13, True, CORAL)


def parrafos(s, items, y=Inches(2.2), x=Inches(0.4), w=Inches(3.7), tam=12.5, h=Inches(2.8)):
    """items: lista de párrafos; cada uno es str o lista de (texto, negrita)."""
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(9)
        p.line_spacing = 1.12
        for tx, neg in ([(it, False)] if isinstance(it, str) else it):
            _run(p, tx, tam, neg)
    return tb


def imagen(s, ruta, x, y, max_w, max_h):
    """Coloca una imagen ajustándola a una caja sin deformarla."""
    with Image.open(ruta) as im:
        iw, ih = im.size
    esc = min(max_w / iw, max_h / ih)
    return s.shapes.add_picture(str(ruta), x, y, width=int(iw * esc), height=int(ih * esc))


def captura_derecha(s, nombre, y=Inches(1.7), max_w=Inches(5.5), max_h=Inches(3.3), x=Inches(4.2)):
    return imagen(s, CAP / f"{nombre}.png", x, y, max_w, max_h)


def pie(s, tx, y=Inches(4.75)):
    tb = s.shapes.add_textbox(Inches(0.4), y, Inches(9.2), Inches(0.35))
    tb.text_frame.word_wrap = True
    _run(tb.text_frame.paragraphs[0], tx, 11, False, GRIS)


# ================================================================ DIAPOSITIVAS
# Se conservan de la presentación actual: 1–11 (incluido tu diagrama oscuro del
# proceso), «// SKILLS // AGENTES // SCRIPTS» con la tríada, «Skill» con la caja
# de herramientas, «AGENTES» y «Agente» con la máquina, y el cierre.

# A · Dónde usar IA y dónde no (patrón de la diapositiva 5 de MujeresEnSEO)
s = nueva("""
TIEMPO 60 s · Va justo después de tu diagrama del proceso.
GUION
Dentro de cada etapa hay dos tipos de trabajo. Uno sale siempre igual: descargar, comparar si un artículo cambió, contar en cuántos canales aparece un tema. Con los mismos datos, el mismo resultado. Ahí no hace falta pensar, hace falta repetir sin equivocarse. Eso va a reglas y a código.
El otro exige mirar y elegir. Cuando digo «tomar una decisión» me refiero a esto: tener varias opciones, mirar la información y elegir una aplicando una regla que podemos escribir. ¿Este tema es nuevo, amplía algo o ya está cubierto? Esa regla escrita, con ejemplos de sí y de no, es el criterio. Y ahí la IA trabaja como clasificador, no como oráculo.
FUENTE: ARQUITECTURA.md §2; MujeresEnSEO, diapositiva 5.
""")
titulo(s, [("Dónde usar IA y dónde no.", False)])
parrafos(s, [
    [("Si el trabajo sale siempre igual: ", True), ("reglas y código. Mismos datos, mismo resultado.", False)],
    [("Si hay que mirar y elegir: ", True), ("IA como clasificador, no como oráculo. Decide contra un criterio escrito.", False)],
    [("Tomar una decisión ", True), ("es elegir entre opciones mirando la información y aplicando una regla que podemos escribir. Esa regla es el criterio.", False)],
], y=Inches(1.7), w=Inches(3.5))
filas = [("", "Sale siempre igual", "Hay que mirar y elegir"),
         ("Revisar lo publicado", "descargar y detectar cambios", "de qué trata y con qué enfoque"),
         ("Encontrar oportunidades", "contar vistas y canales", "¿nuevo, ampliación o cubierto?"),
         ("Investigar", "comprobar que las fuentes responden", "construir el argumento"),
         ("Adaptar", "longitudes, enlaces, formato", "el mensaje de cada canal"),
         ("Aprender", "guardar versiones y estados", "convertir una corrección en regla"),
         ("", "Usar reglas / scripts", "Usar IA con criterio escrito")]
x0, y0 = Inches(4.2), Inches(1.75)
anchos = [Inches(1.55), Inches(1.95), Inches(1.95)]
alto = Inches(0.42)
for i, fila in enumerate(filas):
    x = x0
    for j, celda in enumerate(fila):
        c = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y0 + i * alto, anchos[j], alto)
        c.fill.solid()
        c.fill.fore_color.rgb = BLANCO if i not in (0, 6) else CREMA
        c.line.color.rgb = TINTA
        c.line.width = Pt(0.75)
        c.shadow.inherit = False
        tf = c.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.06)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
        _run(p, celda, 9, i in (0, 6) or j == 0, TINTA)
        x += anchos[j]

# B · De la regla sale la arquitectura (texto aprobado, verbatim de MujeresEnSEO)
s = nueva("""
TIEMPO 40 s
GUION
Y de esa regla sale la arquitectura. No empecé pensando «voy a hacer algo con agentes porque están de moda». Empecé separando el trabajo.
Lo que no quiero que cambie nunca va a scripts: descargar, comparar, contar, comprobar. Lo que es método va a skills: cómo se analiza un artículo, cómo se investiga un tema. Lo que necesita criterio va a agentes: decidir cobertura, verificar afirmaciones, redactar sin inventar.
El sistema no usa IA en todas partes. Usa IA donde el criterio aporta valor y código donde la variabilidad sería un problema.
FUENTE: MujeresEnSEO, diapositiva 6 (principio conservado en el handoff).
""")
titulo(s, [("… de la regla sale la arquitectura.", False)])
parrafos(s, ["Lo repetible va a scripts.", "El método va a skills.", "Lo que requiere criterio va a agentes."],
         y=Inches(1.9), w=Inches(8), tam=20)
pie(s, "Es una regla orientativa: una skill puede incluir scripts, y un agente puede ejecutar operaciones fijas con herramientas.")

# (Aquí van tus diapositivas «// SKILLS // AGENTES // SCRIPTS» y «Skill»)

# C · Cómo se escribe una skill: la description (artículo de Aina, 2026-07-05)
s = nueva("""
TIEMPO 90 s
GUION
Antes de ver la skill del sistema, cómo se escribe una. Una skill casi nunca falla porque alguien no sepa escribir Markdown. Falla antes.
A la izquierda, la versión mala: «Ayuda con reuniones». ¿Tiene que hacer un acta? ¿Una agenda? ¿Tickets? El agente no lo sabe, así que decide por aproximación. Y lo peligroso es que sonará bien aunque haya elegido mal.
A la derecha, la buena. La description no resume la skill: le dice al agente cuándo usarla y cuándo no. Qué entra, qué sale y qué tareas parecidas quedan fuera. Y debajo, un procedimiento con una instrucción que importa más que las demás: «No inventes responsables, fechas ni compromisos».
La prueba es simple: escribe ejemplos en los que la skill debería entrar y ejemplos parecidos en los que debería quedarse fuera. Si no puedes distinguirlos, la skill no está definida.
FUENTE: Think&Hack, «Cómo crear una skill para que tu agente tome mejores decisiones» (2026-07-05).
""")
titulo(s, [("Una skill no falla por el Markdown.", False)])
subtitulo(s, "La description no resume la skill. Le dice al agente cuándo usarla y cuándo no.", w=Inches(9))
imagen(s, CAP / "desc_mala.png", Inches(0.4), Inches(2.15), Inches(4.5), Inches(2.4))
imagen(s, CAP / "desc_buena.png", Inches(5.1), Inches(2.15), Inches(4.5), Inches(2.4))
pie(s, "Prueba: escribe peticiones en las que debería entrar y parecidas en las que debería quedarse fuera. Si no las distingues, no está definida.", y=Inches(4.62))

# D · La skill del sistema (captura real)
s = nueva("""
TIEMPO 60 s
GUION
Esta es la skill real de la primera etapa, tal como está en el repositorio. Fijaos en la description: dice qué hace, descargar y mantener actualizado lo publicado, y cuándo usarla: antes de analizar la biblioteca o tras publicar un artículo nuevo. Son las palabras que yo usaría al pedirlo.
Debajo, qué lee y qué escribe. Nada de la publicación está escrito en el código: otra persona cambia un archivo de configuración y le sirve para su newsletter.
Y una regla que escribí porque me pasó: un error nunca se presenta como «sin novedades». Eso es lo que la IA no puede saber y por eso está escrito.
FUENTE: .agents/skills/recuperar-newsletters/SKILL.md (archivo real).
""")
titulo(s, [("La ", False), ("skill", True), (" guarda el método", False)])
subtitulo(s, "En el sistema · recuperar-newsletters")
parrafos(s, [
    "La description dice qué hace y cuándo usarla, con las palabras que yo usaría al pedirlo.",
    "Qué lee y qué escribe. Nada de mi publicación está en el código: se cambia la configuración.",
    [("La regla que aprendí por las malas: ", True), ("un fallo nunca se presenta como «sin novedades».", False)],
])
captura_derecha(s, "skill_md", max_h=Inches(3.4))

# E · Qué va dentro y qué va fuera (árbol real + práctica)
s = nueva("""
TIEMPO 75 s
GUION
Una skill es una carpeta. Lo que el agente necesita siempre va en SKILL.md: cuándo se usa, los pasos, qué evitar, cómo revisar la salida. Lo que solo necesita a veces va en references: un esquema, unos criterios largos. Las comprobaciones repetibles van en scripts.
Dos consejos que me ahorraron mucho tiempo. Primero: no escribas lo que la IA ya sabe. Escribe lo que no puede saber: tus pasos, lo que siempre sale mal y lo que no quieres ver en el resultado. Segundo: SKILL.md corto, por debajo de quinientas líneas, y el resto en references. Si lo cargas todo cada vez, las instrucciones compiten con la tarea.
Y antes de cerrarla, revisa cada frase con una pregunta: ¿esto cambia alguna decisión del agente? Si le dice cuándo activarse, qué pasos seguir, qué evitar o cómo revisar, se queda. Si suena bien pero no cambia lo que hará, sobra.
FUENTE: artículo 2026-07-05 (estructura y la pregunta); taller de Fred Harper (lo que la IA no sabe, < 500 líneas); .agents/skills/ (árbol real).
""")
titulo(s, [("Qué va dentro y qué va fuera", False)])
subtitulo(s, "Una carpeta, tres papeles")
parrafos(s, [
    [("SKILL.md: ", True), ("lo que necesita siempre. Cuándo se usa, los pasos, qué evitar, cómo revisar. Menos de 500 líneas.", False)],
    [("references/: ", True), ("lo que solo necesita a veces. Un esquema, unos criterios largos.", False)],
    [("scripts/: ", True), ("las comprobaciones repetibles.", False)],
    [("Antes de cerrarla: ", True), ("¿esta frase cambia alguna decisión del agente? Si no, sobra.", False)],
], tam=11.5)
captura_derecha(s, "arbol_skills", max_h=Inches(3.3))

# F · Carga progresiva (estado real de las ocho skills)
s = nueva("""
TIEMPO 60 s
GUION
El sistema tiene ocho skills; tres construidas y cinco especificadas. Si la IA tuviera que leer las ocho enteras cada vez que le pido algo, se confundiría y costaría más. Funciona en tres pasos. Al arrancar solo ve el nombre y la descripción de cada una: una lista corta, unas cien palabras por skill. Cuando pido «actualiza lo publicado», reconoce qué descripción encaja y abre esa skill: lee sus instrucciones. Y solo si la tarea lo necesita, abre los recursos: el programa, el esquema.
Por eso la description es tan importante. Una skill que la IA no reconoce es una skill que no tienes. Y si dos descripciones se solapan, está adivinando. Pídele a la IA que revise tus descripciones: qué hace cada una, cuándo debería activarse y dónde se solapa con otra.
FUENTE: README.md (tabla real); vídeo «Anthropic dejó de crear agentes»; taller de Fred Harper.
""")
titulo(s, [("La IA lee solo lo que necesita", False)])
subtitulo(s, "Ocho skills, tres niveles de lectura")
parrafos(s, [
    [("1 · Siempre: ", True), ("nombre y descripción de cada skill. Una lista corta.", False)],
    [("2 · Si la petición encaja: ", True), ("abre ese SKILL.md y lee las instrucciones.", False)],
    [("3 · Solo si hace falta: ", True), ("los recursos. El programa, el esquema, los ejemplos.", False)],
    [("Si dos descripciones se solapan, ", True), ("la IA está adivinando. Pídele que las revise.", False)],
], tam=11.5)
captura_derecha(s, "estado", max_h=Inches(3.3))
pie(s, "Una skill que la IA no reconoce es una skill que no tienes.")

# G · Script (ilustración de Aina, MujeresEnSEO) + principio DRY
s = nueva("""
TIEMPO 60 s
GUION
Los scripts. Si una operación tiene que salir igual cada vez, no se la pidas a la IA: escríbela en código. Descargar, comparar huellas, contar, comprobar enlaces. El script tiene una virtud muy aburrida y muy valiosa: no improvisa.
Los ingenieros de Anthropic contaban algo que a mí también me pasaba: la IA escribía el mismo programa una y otra vez, cada vez un poco distinto, gastando tokens en redescubrir una solución que ya existía. La solución fue pedirle que lo guardara dentro de la skill, como una herramienta para su yo futuro. La próxima vez no lo reescribe: lo ejecuta.
Así que cuando la IA escriba un programa que te dé un resultado que te gusta, no lo dejes en el chat. Dile: guárdalo en la carpeta scripts de la skill y actualiza las instrucciones para que la próxima vez lo use.
FUENTE: MujeresEnSEO, diapositiva 10; vídeo Anthropic, principio 1.
""")
titulo(s, [("Script", False)])
subtitulo(s, "¿Por qué gastar tokens para que la IA vuelva a aprender una operación que ya puedes fijar en código?", w=Inches(4))
parrafos(s, [
    "Un script fija en código una operación que no debería depender del modelo. Mismo input, mismo output.",
    [("Guarda el programa que funcionó. ", True), ("Si la IA escribe un script que te da el resultado que quieres, pídele que lo guarde dentro de la skill. La próxima vez lo ejecuta, no lo reescribe.", False)],
], y=Inches(2.55), w=Inches(4), tam=11.5)
imagen(s, ILU / "mujeres" / "s10_image12.png", Inches(4.5), Inches(0.5), Inches(5.3), Inches(4.5))

# H · El script en el sistema (informe real)
s = nueva("""
TIEMPO 60 s
GUION
Esto es el informe real de la ejecución de hoy. El script recorrió el archivo de Think&Hack: seis páginas, cincuenta y cinco newsletters, cincuenta nuevas, cinco que ya tenía sin cambios, cero errores. Cada artículo queda con su identificador y su huella: si el texto cambia aunque sea una letra, la huella cambia, y solo se vuelve a analizar lo que cambió.
Ahí no hay IA. Si falla una página, lo dice y marca la ejecución como incompleta. Nunca presenta un fallo como «sin novedades». Esa es la base sobre la que después trabaja el modelo.
FUENTE: biblioteca/substack/control/ejecuciones/20261014T184051Z.json (archivo real).
""")
titulo(s, [("El ", False), ("script", True), (" hace lo repetible", False)])
subtitulo(s, "recuperar-newsletters · hoy, descarga completa")
parrafos(s, [
    [("55 newsletters, 6 páginas, 0 errores. ", True), ("Cada artículo con su identificador y su huella: si cambia una letra, cambia la huella.", False)],
    "Solo se vuelve a analizar lo que ha cambiado.",
    "Si una página falla, la ejecución queda marcada como incompleta. Un fallo nunca se presenta como «sin novedades».",
])
captura_derecha(s, "ejecucion", max_h=Inches(3.4))

# I · Se prueba antes de usarse (tests reales + cómo probar una skill)
s = nueva("""
TIEMPO 60 s
GUION
Y se prueba antes de usarse. Estas son las once pruebas reales del script, con datos simulados: repetir no duplica, un artículo modificado conserva la versión anterior, un extracto no sustituye una copia completa, una página fallida deja la ejecución incompleta. Si una falla, el script no se usa.
Con una skill pasa igual, y hay dos preguntas distintas. ¿Se activa cuando toca? Prueba una petición directa, la misma con otras palabras y una que no tiene nada que ver. ¿Trabaja bien? Lanza la misma petición con y sin skill: si la salida es igual, la skill no aporta. Y dale un caso al que le falte un dato: queremos que lo detecte, no que lo rellene.
FUENTE: test_recuperar_newsletters.py (salida real); artículo 2026-07-05 (con y sin skill); taller de Fred Harper (tres peticiones).
""")
titulo(s, [("Se prueba antes de usarse", False)])
subtitulo(s, "Script: datos simulados. Skill: peticiones.")
parrafos(s, [
    [("¿Se activa cuando toca? ", True), ("Una petición directa, la misma con otras palabras y una que no le toca.", False)],
    [("¿Trabaja bien? ", True), ("La misma petición con y sin skill. Si la salida es igual, no aporta.", False)],
    [("Un caso al que le falte un dato. ", True), ("Queremos que lo detecte, no que lo rellene.", False)],
])
captura_derecha(s, "tests", max_h=Inches(3.4))

# (Aquí van tus diapositivas «AGENTES» y «Agente» con la máquina)

# J · El agente en el sistema (ficha real)
s = nueva("""
TIEMPO 75 s
GUION
Segunda skill, analizar-biblioteca, y aquí entra el modelo. Lee cada artículo completo y rellena una ficha: tema, enfoque, tesis, entidades, nivel. Eso es juicio. Pero fijaos en el reparto: el identificador, la fecha, la huella y las métricas los impone el script desde el inventario. Si el modelo los trae distintos, se ignoran. Y el script valida la ficha contra un esquema antes de guardarla. Si falta un campo, no escribe nada y dice cuál.
El modelo nunca escribe en la biblioteca directamente. Entrega un JSON; el script lo valida y lo guarda.
Hoy hay cincuenta y cinco fichas y doce temas. Esta es la de «Adiós a los GPT personalizados»: tema «Skills para agentes», enfoque «tutorial con prueba propia», y la tesis en una frase. Con esto, la siguiente etapa no tiene que releer el artículo para saber si un tema ya está cubierto.
FUENTE: biblioteca/substack/fichas/217631531.json (real); prompts/02-crear-analizar-biblioteca.md.
""")
titulo(s, [("El modelo juzga, el ", False), ("script", True), (" valida", False)])
subtitulo(s, "analizar-biblioteca · 55 fichas, 12 temas")
parrafos(s, [
    [("El modelo ", True), ("lee el artículo completo y rellena tema, enfoque, tesis, entidades y nivel.", False)],
    [("El script ", True), ("impone id, fecha, huella y métricas desde el inventario, valida contra el esquema y guarda. Si falta un campo, no escribe nada.", False)],
    "El modelo nunca escribe en la biblioteca. Entrega un JSON; el script decide si entra.",
], tam=11.5)
captura_derecha(s, "ficha", max_h=Inches(3.45))

# K · El contrato (esquema real)
s = nueva("""
TIEMPO 45 s
GUION
A esto lo llamo contrato: lo que una etapa se compromete a entregar a la siguiente, siempre con los mismos campos. Este es el esquema real de la ficha. Trece campos obligatorios, valores permitidos para el enfoque y el nivel, entre uno y cinco subtemas.
¿Por qué importa? Porque la siguiente etapa no interpreta un texto distinto cada vez: lee campos. Y porque dos fichas se pueden comparar, que es la base para saber si un tema ya está cubierto. Sin contrato, cada paso reinterpreta el anterior.
FUENTE: .agents/skills/analizar-biblioteca/references/ficha-articulo.schema.json (real).
""")
titulo(s, [("El contrato entre etapas", False)])
subtitulo(s, "Los mismos campos, siempre")
parrafos(s, [
    "Lo que una etapa se compromete a entregar a la siguiente, con campos fijos y valores permitidos.",
    "La siguiente etapa no interpreta un texto distinto cada vez: lee campos.",
    "Dos fichas se pueden comparar. Así se sabe si un tema ya está cubierto.",
])
captura_derecha(s, "schema", max_h=Inches(3.45))

# L0 · Interés real: lo calcula el script (YouTube real)
s = nueva("""
TIEMPO 60 s
GUION
Antes de que el agente decida nada, un script ha hecho los números. Esto es un vídeo real de los ciento cincuenta que el sistema recogió hoy de cinco canales. Trescientas veintinueve mil visualizaciones. ¿Es mucho? Depende del canal. Por eso el script calcula el múltiplo de la mediana: las vistas de este vídeo divididas entre lo que suele hacer ese canal. Siete veces su mediana: interés alto. Cincuenta mil vistas pueden ser mucho en un canal pequeño y nada en uno grande; el múltiplo corrige eso.
Esto no se lo pido a la IA. Es una división. Con los mismos datos, el mismo resultado, y queda guardado con la fecha de captura. Sobre esta base, después, el agente lee y juzga.
FUENTE: biblioteca/youtube/videos/AYQtRqW1xX4/meta.json (real); ARQUITECTURA.md §8.2.
""")
titulo(s, [("Interés real: lo calcula el ", False), ("script", True)])
subtitulo(s, "explorar-youtube · 150 vídeos, 5 canales")
parrafos(s, [
    [("Múltiplo de la mediana: ", True), ("las vistas de este vídeo divididas entre lo que suele hacer su canal. Corrige el tamaño del canal.", False)],
    [("7 veces su mediana: interés alto. ", True), ("Es una división, no un juicio. Mismos datos, mismo resultado.", False)],
    "Sobre esta base, después, el agente lee y decide.",
])
captura_derecha(s, "youtube", max_h=Inches(3.45))

# L · Qué es una oportunidad (las cuatro condiciones, texto real)
s = nueva("""
TIEMPO 75 s
GUION
Tercera etapa: decidir si un tema es una oportunidad. Aquí hay que mirar y elegir, y el reparto es el mismo. Un tema es oportunidad solo si cumple cuatro condiciones. Dos las calcula un script con números: interés real, vistas frente a la mediana del propio canal, y presencia en dos o más canales. Dos las juzga el agente después de leer los artículos parecidos completos: que no esté cubierto y que podamos aportar algo concreto. Sin aportación es copiar, no una oportunidad.
Y la decisión final no la toma el agente. La toma el script aplicando las cuatro condiciones al juicio del agente. No hay una nota única: cada condición queda visible y la persona elige con los motivos delante.
Un buen agente tiene una fase concreta. «Hazme una estrategia de contenidos» no es una fase. «Decide si este tema está cubierto, leyendo estos cinco artículos» sí.
FUENTE: ARQUITECTURA.md §8.4 (texto real); prompts/04-crear-detectar-oportunidades.md.
""")
titulo(s, [("El ", False), ("agente", True), (" decide dentro de límites", False)])
subtitulo(s, "detectar-oportunidades · las cuatro condiciones")
parrafos(s, [
    [("Dos condiciones las calcula el script: ", True), ("interés real y presencia en varios canales.", False)],
    [("Dos las juzga el agente ", True), ("tras leer los artículos parecidos: no está cubierto y podemos aportar algo.", False)],
    [("La decisión final la aplica el script ", True), ("con las cuatro condiciones. Cada una queda visible. La persona elige.", False)],
], tam=11.5)
captura_derecha(s, "condiciones", max_h=Inches(3.45))

# M · El arnés: reglas reales + evidencia
s = nueva("""
TIEMPO 90 s
GUION
Nos falta una palabra: arnés. La skill conserva cómo se hace una tarea. El agente recibe el trabajo y decide cómo avanzar. El arnés controla el entorno en el que trabaja y qué cuenta como terminado. Las skills no se activan unas a otras por estar en la misma carpeta: alguien comprueba que la entrega anterior existe antes de pasar a la siguiente etapa.
Esto es texto real del sistema, de la etapa de investigación. Cada afirmación que va a entrar en el texto queda verificada, refutada o sin confirmar, con su fuente. Y una regla que me gusta mucho: una afirmación verificada necesita al menos una fuente que no sea el propio vídeo de origen. El vídeo no basta para verificarse a sí mismo. Si una pieza dice «este método reduce un 80 % el tiempo» y ninguna fuente lo sostiene, no entra.
Y abajo: nunca se sobrescribe una investigación aprobada. La aprobación la marca una persona cambiando el estado. Una fase no está hecha porque la IA lo diga. Está hecha si su archivo existe con los campos del contrato. Eso es evidencia: una prueba que existe fuera de la respuesta de la IA. Que la IA relea su texto y diga «me parece bien» no lo es.
FUENTE: prompts/05-crear-investigar-tema.md (texto real); Think&Hack, «Cómo construir un arnés» (2026-07-19); MujeresEnSEO, «El arnés (Default-FAIL)».
""")
titulo(s, [("El ", False), ("arnés", True), (": qué cuenta como terminado", False)])
subtitulo(s, "Skill: conserva el método · Agente: decide · Arnés: controla el entorno y el cierre", w=Inches(9))
parrafos(s, [
    [("Una fase no está hecha porque la IA lo diga. ", True), ("Está hecha si su archivo existe con los campos del contrato.", False)],
    [("Evidencia: ", True), ("una prueba fuera de la respuesta de la IA. El vídeo de origen no basta para verificarse a sí mismo.", False)],
    "Nunca se sobrescribe una investigación aprobada. La aprueba una persona.",
], y=Inches(2.25))
captura_derecha(s, "arnes_reglas", y=Inches(2.2), max_h=Inches(2.9))

# N · Aprender: cómo se diagnostica (texto real del prompt 08 + práctica)
s = nueva("""
TIEMPO 75 s
GUION
Última etapa, y la segunda restricción del objetivo: no perder lo aprendido. Cada vez que corriges a la IA en un chat y cierras la ventana, esa lección se pierde. Y quiero ser precisa: una skill tampoco aprende sola. Lo que hace el sistema es convertir la corrección en una regla escrita.
Esta es la guía de estilo real, generada hoy a partir de los cincuenta y cinco artículos, y abajo la sección que importa: reglas aprendidas. La última skill compara el borrador que escribió la IA con lo que publicamos, párrafo a párrafo. La IA resume las diferencias como reglas, cada una con un ejemplo antes y después, y se añaden a la guía de estilo. La próxima redacción la lee antes de empezar. No añade una regla repetida. Y un descarte sin motivo no enseña nada: si descarto una oportunidad, escribo por qué, y la siguiente búsqueda lo consulta.
La práctica para vuestras skills es esta. Cuando algo sale mal, antes de arreglar la respuesta, decidid qué falló: ¿el proceso, el contexto, una regla débil o un programa poco fiable? Cada causa se arregla en un sitio distinto: las instrucciones, un archivo de referencia, una regla explícita o el script. Actualizad la skill en el lugar permanente más pequeño, volved a ejecutar la tarea y comprobad. Y hacedlo el día que pasa.
FUENTE: biblioteca/estilo.md (generada hoy); prompts/08-crear-registrar-resultado.md; vídeo Anthropic, principio 3; taller de Fred Harper.
""")
titulo(s, [("Arregla la ", False), ("skill", True), (", no la respuesta", False)])
subtitulo(s, "registrar-resultado · de corrección a regla")
parrafos(s, [
    [("Una skill no aprende sola. ", True), ("La corrección se escribe como regla, con un ejemplo antes y después. La próxima redacción la lee.", False)],
    [("¿Qué falló? ", True), ("El proceso, el contexto, una regla débil o un programa poco fiable. Cada causa se arregla en un sitio distinto.", False)],
    "Actualiza la skill en el lugar permanente más pequeño, vuelve a ejecutar y comprueba. Hazlo el día que pasa.",
], tam=11.5)
captura_derecha(s, "estilo", max_h=Inches(3.45))

# O · El resultado hoy (mapa real)
s = nueva("""
TIEMPO 45 s
GUION
Esto es lo que el sistema ha producido hoy, con tres de las ocho skills construidas. Cincuenta y cinco newsletters descargadas, cincuenta y cinco fichas, un mapa de temas y una guía de estilo generados: once artículos sobre agentes, seis sobre estrategia, seis sobre escribir con IA, cuatro sobre skills. Y ciento cincuenta vídeos de cinco canales con sus señales calculadas. Es la respuesta a la primera pregunta del proceso: qué hemos tratado y desde qué perspectiva, y qué se está publicando fuera.
Las otras cinco skills están especificadas en el repositorio como prompts que cualquiera puede ejecutar en orden sobre una carpeta vacía. Eso es lo que comparto.
FUENTE: biblioteca/mapa-temas.md (generado hoy); README.md.
""")
titulo(s, [("El resultado, hoy", False)])
subtitulo(s, "Tres skills construidas de ocho")
parrafos(s, [
    [("55 newsletters. 55 fichas. 12 temas. Una guía de estilo. ", True), ("Qué hemos tratado, cuándo y desde qué enfoque.", False)],
    [("150 vídeos de 5 canales ", True), ("con sus señales calculadas: qué se está publicando fuera.", False)],
    "Las cinco skills restantes están escritas como prompts que cualquiera ejecuta en orden sobre una carpeta vacía.",
])
captura_derecha(s, "mapa_temas", max_h=Inches(3.45))

# P · Cinco preguntas antes de escribir una skill (artículo de Aina)
s = nueva("""
TIEMPO 60 s
GUION
Para empezar mañana no hace falta montar las cinco etapas. Hace falta una tarea que se repita y que sepáis revisar. Y antes de escribir la skill, cinco preguntas.
¿Qué tarea repetida quieres mejorar? ¿Qué hace mal el agente sin skill? ¿Qué criterio experto debe aplicar? ¿Cuándo no debe activarse? ¿Qué error debe evitar? Si podéis contestar las cinco, tenéis la skill. Si no podéis contestar la segunda, todavía no la necesitáis.
Después, la prueba incómoda: la misma petición con y sin skill. Una skill empieza a tener sentido cuando evita un error que antes ocurría. Eso ya empieza a ser sistema.
FUENTE: Think&Hack, 2026-07-05 («Cinco preguntas antes de escribir una skill»).
""")
titulo(s, [("Cinco preguntas antes de una skill", False)])
preguntas = [("¿Qué tarea repetida quieres mejorar?", "Una. Que sepas revisar."),
             ("¿Qué hace mal el agente sin skill?", "Si no lo sabes, aún no la necesitas."),
             ("¿Qué criterio experto debe aplicar?", "Lo que tú miras y la IA no sabe."),
             ("¿Cuándo no debe activarse?", "Las tareas parecidas que quedan fuera."),
             ("¿Qué error debe evitar?", "El que ya te ha pasado.")]
y = Inches(1.75)
for i, (q, a) in enumerate(preguntas, 1):
    c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.45), y, Inches(0.42), Inches(0.42))
    c.fill.solid(); c.fill.fore_color.rgb = CORAL; c.line.fill.background(); c.shadow.inherit = False
    c.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    c.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    _run(c.text_frame.paragraphs[0], str(i), 13, True, BLANCO)
    tb = s.shapes.add_textbox(Inches(1.05), y - Inches(0.04), Inches(8.5), Inches(0.5))
    tb.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tb.text_frame.paragraphs[0]
    _run(p, q + "  ", 15, True)
    _run(p, a, 12.5, False, GRIS)
    y += Inches(0.58)
pie(s, "Después, la prueba incómoda: la misma petición con y sin skill. Si la salida es igual, la skill no aporta.")

prs.save(SALIDA)
print("Guardado", SALIDA, "·", len(prs.slides), "diapositivas")
